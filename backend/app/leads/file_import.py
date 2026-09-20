from __future__ import annotations

import csv
import io
import re
from pathlib import Path
from urllib.parse import urlparse

from openpyxl import load_workbook

from app.leads.parser import ParsedLead, PUBLIC_EMAIL_DOMAINS, valid_phone

MAX_ROWS = 5000
SUPPORTED_SUFFIXES = {".csv", ".xlsx"}


def _norm_header(value: object) -> str:
    text = str(value or "").strip().lower()
    text = re.sub(r"[^a-z0-9]+", " ", text)
    return " ".join(text.split())


ALIASES: dict[str, tuple[str, ...]] = {
    "company_name": ("company", "company name", "business", "business name", "organization", "organisation"),
    "first_name": ("first name", "firstname", "given name"),
    "last_name": ("last name", "lastname", "surname", "family name"),
    "full_name": ("name", "contact", "contact name", "full name", "decision maker"),
    "job_title": ("job title", "title", "role", "position"),
    "email": ("email", "email address", "work email", "business email"),
    "email_status": ("email status", "verification", "verification status", "email verification"),
    "phone": ("phone", "phone number", "mobile", "telephone", "tel"),
    "website": ("website", "website url", "url", "company website", "domain"),
    "linkedin_url": ("linkedin", "linkedin url", "linkedin profile"),
    "instagram_url": ("instagram", "instagram url", "instagram profile"),
    "facebook_url": ("facebook", "facebook url", "facebook profile"),
    "city": ("city", "town"),
    "region": ("region", "state", "province", "state region"),
    "country": ("country",),
    "score": ("score", "lead score", "fit score"),
}


def _column_map(headers: list[str]) -> dict[str, int]:
    normalized = [_norm_header(header) for header in headers]
    result: dict[str, int] = {}
    for field, aliases in ALIASES.items():
        for alias in aliases:
            if alias in normalized:
                result[field] = normalized.index(alias)
                break
    return result


def _value(row: list[object], mapping: dict[str, int], field: str) -> str | None:
    index = mapping.get(field)
    if index is None or index >= len(row):
        return None
    raw = row[index]
    if raw is None:
        return None
    text = str(raw).strip()
    return text or None


def _normalize_url(value: str | None) -> str | None:
    if not value:
        return None
    text = value.strip()
    if not text:
        return None
    if text.lower().startswith(("http://", "https://")):
        return text
    if "." in text and " " not in text:
        return f"https://{text}"
    return text


def _domain(website: str | None, email: str | None) -> str | None:
    if website:
        try:
            host = (urlparse(website).hostname or "").lower().removeprefix("www.")
            if host:
                return host
        except ValueError:
            pass
    if email and "@" in email:
        candidate = email.rsplit("@", 1)[1].lower().strip()
        if candidate not in PUBLIC_EMAIL_DOMAINS:
            return candidate
    return None


def _split_name(value: str | None) -> tuple[str | None, str | None]:
    if not value:
        return None, None
    parts = value.strip().split()
    if not parts:
        return None, None
    if len(parts) == 1:
        return parts[0], None
    return parts[0], " ".join(parts[1:])


def _score(value: str | None) -> float | None:
    if not value:
        return None
    try:
        score = float(str(value).replace("%", "").strip())
    except ValueError:
        return None
    if score < 0:
        return 0.0
    if score > 100:
        return 100.0
    return score


def _email_status(email: str | None, value: str | None) -> str | None:
    if not email:
        return None
    normalized = (value or "").strip().lower()
    if normalized in {"verified", "valid", "deliverable", "confirmed"}:
        return "verified"
    if normalized in {"invalid", "undeliverable", "bad"}:
        return "invalid"
    return "unverified"


def _rows_from_csv(payload: bytes) -> list[list[object]]:
    text: str
    try:
        text = payload.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = payload.decode("latin-1")
    return [list(row) for row in csv.reader(io.StringIO(text))]


def _rows_from_xlsx(payload: bytes) -> list[list[object]]:
    workbook = load_workbook(io.BytesIO(payload), read_only=True, data_only=True)
    sheet = workbook.active
    rows = [list(row) for row in sheet.iter_rows(values_only=True)]
    workbook.close()
    return rows


def parse_lead_file(filename: str, payload: bytes) -> tuple[list[ParsedLead], dict[str, str]]:
    suffix = Path(filename or "").suffix.lower()
    if suffix not in SUPPORTED_SUFFIXES:
        raise ValueError("Upload a .xlsx or .csv file.")
    if not payload:
        raise ValueError("The uploaded file is empty.")

    rows = _rows_from_xlsx(payload) if suffix == ".xlsx" else _rows_from_csv(payload)
    rows = [row for row in rows if any(str(cell or "").strip() for cell in row)]
    if len(rows) < 2:
        raise ValueError("The sheet needs a header row and at least one lead row.")
    if len(rows) - 1 > MAX_ROWS:
        raise ValueError(f"Lead-sheet imports are limited to {MAX_ROWS} rows per file.")

    headers = [str(cell or "").strip() for cell in rows[0]]
    mapping = _column_map(headers)
    if not any(field in mapping for field in ("company_name", "email", "website", "phone", "full_name", "first_name")):
        raise ValueError("No recognizable lead columns were found. Include headers such as Company, Email, Website, Phone, or Contact Name.")

    parsed: list[ParsedLead] = []
    for row in rows[1:]:
        company = _value(row, mapping, "company_name")
        email = (_value(row, mapping, "email") or "").strip().lower() or None
        website = _normalize_url(_value(row, mapping, "website"))
        phone = valid_phone(_value(row, mapping, "phone"))
        first = _value(row, mapping, "first_name")
        last = _value(row, mapping, "last_name")
        if not first and not last:
            first, last = _split_name(_value(row, mapping, "full_name"))

        city = _value(row, mapping, "city")
        region = _value(row, mapping, "region")
        country = _value(row, mapping, "country")
        location = ", ".join(part for part in (city, region, country) if part) or None

        lead = ParsedLead(
            company_name=company,
            website=website,
            domain=_domain(website, email),
            first_name=first,
            last_name=last,
            job_title=_value(row, mapping, "job_title"),
            email=email,
            email_status=_email_status(email, _value(row, mapping, "email_status")),
            phone=phone,
            linkedin_url=_normalize_url(_value(row, mapping, "linkedin_url")),
            instagram_url=_normalize_url(_value(row, mapping, "instagram_url")),
            facebook_url=_normalize_url(_value(row, mapping, "facebook_url")),
            region=location,
            source="sheet_import",
            source_url=website,
            source_query=f"Imported from {filename}",
            score=_score(_value(row, mapping, "score")),
        )
        if any((lead.company_name, lead.email, lead.website, lead.phone, lead.linkedin_url, lead.instagram_url, lead.facebook_url)):
            parsed.append(lead)

    detected = {field: headers[index] for field, index in mapping.items() if index < len(headers)}
    return parsed, detected
